"""Controller layer: handle commands and orchestrate Model + View."""

import logging
import re
from datetime import datetime

from models.todo import Todo, TodoModel
from models.user_prefs import UserPrefs
from views.todo_view import TodoView
from platforms.base import BotPlatform, Update
from i18n import t

logger = logging.getLogger(__name__)


class TodoController:

    def __init__(self, platform: BotPlatform):
        self.platform = platform
        self.view = TodoView()
        self.model = TodoModel()
        self._commands = {
            "/start": self._cmd_start,
            "/help": self._cmd_help,
            "/add": self._cmd_add,
            "/list": self._cmd_list,
            "/all": self._cmd_all,
            "/done": self._cmd_done,
            "/undone": self._cmd_undone,
            "/del": self._cmd_delete,
            "/delete": self._cmd_delete,
            "/edit": self._cmd_edit,
            "/note": self._cmd_note,
            "/priority": self._cmd_priority,
            "/due": self._cmd_due,
            "/setdue": self._cmd_setdue,
            "/repeat": self._cmd_repeat,
            "/search": self._cmd_search,
            "/assign": self._cmd_assign,
            "/clear": self._cmd_clear,
            "/stats": self._cmd_stats,
            "/detail": self._cmd_detail,
            "/export": self._cmd_export,
            "/lang": self._cmd_lang,
        }

    def _get_lang(self, update: Update) -> str:
        prefs = UserPrefs.get(update.user.id)
        lang = prefs.get("language", "auto")
        if lang == "auto":
            lang = UserPrefs.detect_language(update.message.text)
        return lang

    def handle_update(self, update: Update):
        """Route incoming update to the appropriate command handler."""
        text = update.message.text.strip()
        if not text:
            return

        self.platform.send_typing(update.chat.id, update.chat.type)

        lang = self._get_lang(update)

        # Check first-use onboarding
        if UserPrefs.is_first_use(update.user.id):
            UserPrefs.mark_welcomed(update.user.id)
            if text == "/start" or not text.startswith("/"):
                reply_text, msg_type = self.view.onboarding(lang)
                self.platform.send_message(
                    chat_id=update.chat.id, text=reply_text,
                    chat_type=update.chat.type, msg_type=msg_type,
                )
                return

        # Parse command
        parts = text.split(maxsplit=1)
        cmd = parts[0].lower()
        args = parts[1] if len(parts) > 1 else ""

        if "@" in cmd:
            cmd = cmd.split("@")[0]

        handler = self._commands.get(cmd)
        if handler:
            reply_text, msg_type = handler(update, args, lang)
        else:
            reply_text, msg_type = t("err_unknown_cmd", lang), 1

        self.platform.send_message(
            chat_id=update.chat.id,
            text=reply_text,
            chat_type=update.chat.type,
            msg_type=msg_type,
            reply_to=update.message.id,
        )

    # --- Commands ---

    def _cmd_start(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        return self.view.welcome(lang)

    def _cmd_help(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        return self.view.help_message(lang)

    def _cmd_add(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        if not args:
            return self.view.error(t("err_add_usage", lang))

        title = args
        priority = 2
        due_date = None
        tags = []

        # Extract #tags
        tag_matches = re.findall(r'#(\S+)', title)
        priority_words = {"high", "medium", "low"}
        for tag in tag_matches:
            low = tag.lower()
            if low in priority_words:
                priority = {"high": 1, "medium": 2, "low": 3}[low]
            else:
                tags.append(tag)
            title = title.replace(f"#{tag}", "").strip()

        # Extract @YYYY-MM-DD due date
        due_match = re.search(r'@(\d{4}-\d{2}-\d{2})', title)
        if due_match:
            due_date = due_match.group(1)
            try:
                datetime.strptime(due_date, "%Y-%m-%d")
            except ValueError:
                return self.view.error(t("err_date_invalid", lang))
            title = title.replace(due_match.group(0), "").strip()

        if not title:
            return self.view.error(t("err_empty_title", lang))

        todo = self.model.create(
            user_id=update.user.id,
            chat_id=update.chat.id,
            chat_type=update.chat.type,
            title=title,
            priority=priority,
            due_date=due_date,
            tags=",".join(tags) if tags else "",
        )
        return self.view.todo_created(todo, lang)

    def _cmd_list(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        args = args.strip().lower()
        if args == "overdue":
            todos = self.model.get_overdue(update.user.id)
            return self.view.todo_list(todos, t("overdue_todos", lang), lang)
        if args.startswith("#"):
            tag = args[1:]
            todos = self.model.list_by_user(update.user.id, status=0, tag=tag)
            return self.view.todo_list(todos, f"#{tag}", lang)
        if args in ("high", "1"):
            todos = self.model.list_by_user(update.user.id, status=0, priority=1)
            return self.view.todo_list(todos, "High Priority", lang)
        if args in ("medium", "2"):
            todos = self.model.list_by_user(update.user.id, status=0, priority=2)
            return self.view.todo_list(todos, "Medium Priority", lang)
        if args in ("low", "3"):
            todos = self.model.list_by_user(update.user.id, status=0, priority=3)
            return self.view.todo_list(todos, "Low Priority", lang)

        todos = self.model.list_by_user(update.user.id, status=0)
        return self.view.todo_list(todos, t("pending_todos", lang), lang)

    def _cmd_all(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todos = self.model.list_by_user(update.user.id)
        return self.view.todo_list(todos, t("all_todos", lang), lang)

    def _cmd_done(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todo_id = self._parse_id(args)
        if not todo_id:
            return self.view.error(t("err_done_usage", lang))
        completed_todo = self.model.complete(todo_id, update.user.id)
        if completed_todo:
            # Handle recurring: recreate if repeat_rule is set
            if completed_todo.repeat_rule:
                self._recreate_recurring(completed_todo)
            return self.view.todo_completed(todo_id, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_undone(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todo_id = self._parse_id(args)
        if not todo_id:
            return self.view.error(t("err_undone_usage", lang))
        if self.model.uncomplete(todo_id, update.user.id):
            return self.view.todo_uncompleted(todo_id, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_delete(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todo_id = self._parse_id(args)
        if not todo_id:
            return self.view.error(t("err_del_usage", lang))
        if self.model.delete(todo_id, update.user.id):
            return self.view.todo_deleted(todo_id, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_edit(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        parts = args.split(maxsplit=1)
        if len(parts) != 2:
            return self.view.error(t("err_edit_usage", lang))
        todo_id = self._parse_id(parts[0])
        new_title = parts[1].strip()
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        if not new_title:
            return self.view.error(t("err_empty_title", lang))
        if self.model.update_title(todo_id, update.user.id, new_title):
            return self.view.title_edited(todo_id, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_note(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        parts = args.split(maxsplit=1)
        if len(parts) != 2:
            return self.view.error(t("err_note_usage", lang))
        todo_id = self._parse_id(parts[0])
        text = parts[1].strip()
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        if self.model.update_description(todo_id, update.user.id, text):
            return self.view.note_added(todo_id, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_priority(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        parts = args.split()
        if len(parts) != 2:
            return self.view.error(t("err_priority_usage", lang))
        todo_id = self._parse_id(parts[0])
        try:
            priority = int(parts[1])
        except ValueError:
            return self.view.error(t("err_priority_invalid", lang))
        if priority not in (1, 2, 3):
            return self.view.error(t("err_priority_invalid", lang))
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        if self.model.update_priority(todo_id, update.user.id, priority):
            return self.view.priority_updated(todo_id, priority, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_due(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        parts = args.split()
        if len(parts) != 2:
            return self.view.error(t("err_due_usage", lang))
        todo_id = self._parse_id(parts[0])
        due_date = parts[1]
        try:
            datetime.strptime(due_date, "%Y-%m-%d")
        except ValueError:
            return self.view.error(t("err_date_invalid", lang))
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        if self.model.update_due_date(todo_id, update.user.id, due_date):
            return self.view.due_date_updated(todo_id, due_date, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_setdue(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todo_id = self._parse_id(args)
        if not todo_id:
            return self.view.error(t("err_setdue_usage", lang))
        todo = self.model.get_by_id(todo_id, update.user.id)
        if not todo:
            return self.view.not_found(todo_id, lang)
        return self.view.set_due_prompt(todo, lang)

    def _cmd_repeat(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        parts = args.split()
        if len(parts) != 2:
            return self.view.error(t("err_repeat_usage", lang))
        todo_id = self._parse_id(parts[0])
        rule = parts[1].lower()
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        if rule == "off":
            rule = ""
        elif rule not in ("daily", "weekly", "monthly"):
            return self.view.error(t("err_repeat_invalid", lang))
        if self.model.update_repeat_rule(todo_id, update.user.id, rule):
            return self.view.repeat_set(todo_id, rule, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_search(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        keyword = args.strip()
        if not keyword:
            return self.view.error(t("err_search_usage", lang))
        todos = self.model.search(update.user.id, keyword)
        return self.view.search_results(todos, keyword, lang)

    def _cmd_assign(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        if update.chat.type == "private":
            return self.view.error(t("err_assign_private", lang))
        parts = args.split()
        if len(parts) < 2:
            return self.view.error(t("err_assign_usage", lang))
        todo_id = self._parse_id(parts[0])
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        assignee_id = self._parse_id(parts[1])
        if not assignee_id:
            return self.view.error(t("err_assign_usage", lang))
        assignee_name = parts[2] if len(parts) > 2 else str(assignee_id)
        if self.model.assign(todo_id, update.user.id, assignee_id, assignee_name):
            return self.view.assigned(todo_id, assignee_name, lang)
        return self.view.not_found(todo_id, lang)

    def _cmd_clear(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        count = self.model.clear_completed(update.user.id)
        return self.view.cleared(count, lang)

    def _cmd_stats(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        counts = self.model.count_by_user(update.user.id)
        return self.view.stats(counts, lang)

    def _cmd_detail(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todo_id = self._parse_id(args)
        if not todo_id:
            return self.view.error(t("err_id_invalid", lang))
        todo = self.model.get_by_id(todo_id, update.user.id)
        if not todo:
            return self.view.not_found(todo_id, lang)
        return self.view.todo_detail(todo, lang)

    def _cmd_export(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        todos = self.model.list_by_user(update.user.id)
        if not todos:
            return self.view.error(t("no_todos", lang))
        filepath = self._generate_excel(todos, update.user.id, lang)
        return t("export_done", lang, path=filepath), 1

    def _cmd_lang(self, update: Update, args: str, lang: str) -> tuple[str, int]:
        new_lang = args.strip().lower()
        if not new_lang:
            return self.view.lang_select(lang)
        if new_lang == "auto":
            UserPrefs.set_language(update.user.id, "auto")
            return self.view.lang_set(lang)
        if new_lang not in ("en", "zh"):
            return self.view.lang_select(lang)
        UserPrefs.set_language(update.user.id, new_lang)
        return self.view.lang_set(new_lang)

    # --- Helpers ---

    def _generate_excel(self, todos: list[Todo], user_id: int, lang: str) -> str:
        """Generate an Excel file and return the file path."""
        import os
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from i18n import priority_label as p_label

        wb = Workbook()
        ws = wb.active
        ws.title = t("export_title", lang)

        # Headers
        headers = ["ID", t("status", lang), t("priority", lang), "Title",
                    t("due_date", lang), t("tags", lang), t("repeat", lang),
                    t("assignee", lang), t("description", lang)]
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        # Data rows
        overdue_fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
        done_font = Font(strikethrough=True, color="999999")
        from datetime import datetime as dt
        today = dt.now().strftime("%Y-%m-%d")

        for row_idx, todo in enumerate(todos, 2):
            status_str = t("completed", lang) if todo.status == 1 else t("pending", lang)
            is_overdue = todo.status == 0 and todo.due_date and str(todo.due_date) < today
            if is_overdue:
                status_str = t("overdue", lang)

            values = [
                todo.id,
                status_str,
                p_label(todo.priority, lang),
                todo.title,
                str(todo.due_date) if todo.due_date else "",
                todo.tags.replace(",", ", ") if todo.tags else "",
                todo.repeat_rule or "",
                todo.assignee_name or "",
                todo.description or "",
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = thin_border
                if todo.status == 1:
                    cell.font = done_font
                if is_overdue:
                    cell.fill = overdue_fill

        # Auto-width
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        # Save
        export_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "exports")
        os.makedirs(export_dir, exist_ok=True)
        filename = f"todos_{user_id}_{int(__import__('time').time())}.xlsx"
        filepath = os.path.join(export_dir, filename)
        wb.save(filepath)
        return os.path.abspath(filepath)

    def _recreate_recurring(self, todo: Todo):
        """When a recurring task is completed, create the next occurrence."""
        from datetime import timedelta
        next_due = None
        if todo.due_date:
            try:
                due = datetime.strptime(str(todo.due_date), "%Y-%m-%d")
                if todo.repeat_rule == "daily":
                    next_due = (due + timedelta(days=1)).strftime("%Y-%m-%d")
                elif todo.repeat_rule == "weekly":
                    next_due = (due + timedelta(weeks=1)).strftime("%Y-%m-%d")
                elif todo.repeat_rule == "monthly":
                    month = due.month % 12 + 1
                    year = due.year + (1 if due.month == 12 else 0)
                    next_due = due.replace(year=year, month=month).strftime("%Y-%m-%d")
            except ValueError:
                pass

        self.model.create(
            user_id=todo.user_id,
            chat_id=todo.chat_id,
            chat_type=todo.chat_type,
            title=todo.title,
            priority=todo.priority,
            due_date=next_due,
            description=todo.description,
            tags=todo.tags,
            repeat_rule=todo.repeat_rule,
            assignee_id=todo.assignee_id,
            assignee_name=todo.assignee_name,
        )

    @staticmethod
    def _parse_id(text: str) -> int | None:
        text = text.strip().lstrip("#")
        try:
            val = int(text)
            return val if val > 0 else None
        except (ValueError, TypeError):
            return None
